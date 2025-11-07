import os
import json
import asyncio
import uvicorn
import socketio
import time
import base64  # For decoding
import io      # To handle binary data in memory
import csv     # For CSV parsing
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from typing import AsyncGenerator, Any, Dict, Optional

# --- Import parsing libraries ---
import pypdf
import docx
import openpyxl
import pptx

# Import client files
from src.graph import build_graph, AgentState
from langchain_core.messages import HumanMessage, AIMessageChunk, ToolMessage
from langgraph.graph import StateGraph
from dotenv import load_dotenv
from langchain_mcp_adapters.client import MultiServerMCPClient

load_dotenv()

# --- Globals to hold the compiled graph and client ---
compiled_graph = None
client = None
FILE_PARSE_TIMEOUT = 60.0

# -----------------------------------------------------------------
# 1. DEBUG-ENABLED STREAMING LOGIC (Unchanged)
# -----------------------------------------------------------------
async def _stream_graph_logic(
    input_text: str, 
    graph: StateGraph, 
    config: dict
) -> AsyncGenerator[tuple[str, str], None]:
    """
    Internal streaming logic that yields structured events.
    """
    
    try:
        current_state = graph.get_state(config) 
        if current_state and current_state.values and "messages" in current_state.values:
            old_message_ids = {m.id for m in current_state.values["messages"]}
        else:
            old_message_ids = set()
    except Exception as e:
        print(f"[WARN] No prior state or error getting state: {e}")
        old_message_ids = set()
    
    input_state = AgentState(messages=[HumanMessage(content=input_text)])
    
    async for message_chunk, metadata in graph.astream(
        input=input_state,
        stream_mode="messages",
        config=config
    ):
        # Check for AIMessageChunk (the LLM's response)
        if isinstance(message_chunk, AIMessageChunk):
            if message_chunk.response_metadata:
                if message_chunk.response_metadata.get("finish_reason") == "tool_calls":
                    continue  # Skip the tool call marker

            if message_chunk.content and message_chunk.content.strip():
                # Yield an "ai" event for the frontend
                yield ("ai", message_chunk.content)

        # Check for ToolMessage (the result from a tool)
        elif isinstance(message_chunk, ToolMessage):
            is_old_memory = message_chunk.id in old_message_ids
            if is_old_memory:
                debug_msg = f"[DEBUG-MEMORY] Replaying tool: {message_chunk.name}...\n"
                print(f"[DEBUG-MEMORY] Tool {message_chunk.name} replayed: {message_chunk.content[:100]}...")
            else:
                debug_msg = f"[DEBUG-NEW] Running tool: {message_chunk.name}...\n"
                print(f"[DEBUG-NEW] Tool {message_chunk.name} returned: {message_chunk.content[:200]}...")
                # --- Yield a separate event for the frontend ---
                yield ("tool_call", message_chunk.name)
            
            # Yield a "debug" event for the terminal
            yield ("debug", debug_msg)

# -----------------------------------------------------------------
# 2. FILE PARSING HELPER (with robust PDF handling)
# -----------------------------------------------------------------
def parse_file_content(file_name: str, base64_data: str) -> str:
    """
    Decodes Base64 data and extracts text based on file extension.
    This runs in a separate thread to avoid blocking asyncio.
    """
    try:
        binary_data = base64.b64decode(base64_data)
        file_stream = io.BytesIO(binary_data)
        _, extension = os.path.splitext(file_name.lower())
        
        text_content = []

        # --- MODIFIED: Added specific try/except for pypdf ---
        if extension == '.pdf':
            try:
                reader = pypdf.PdfReader(file_stream)
                for page in reader.pages:
                    text_content.append(page.extract_text() or "") # Add 'or ""' for blank pages
                return "\n".join(text_content)
            except Exception as pdf_error:
                print(f"[ERROR] pypdf failed to parse {file_name}: {pdf_error}")
                return f"[Error: Failed to parse PDF file '{file_name}'. It may be corrupted, password-protected, or have an unsupported format.]"
        # --- END MODIFIED ---

        elif extension == '.docx':
            doc = docx.Document(file_stream)
            for para in doc.paragraphs:
                text_content.append(para.text)
            return "\n".join(text_content)

        elif extension == '.xlsx':
            wb = openpyxl.load_workbook(file_stream, read_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_content.append(f"--- Sheet: {sheet_name} ---")
                for row in sheet.iter_rows():
                    row_text = [str(cell.value) if cell.value is not None else "" for cell in row]
                    text_content.append(", ".join(row_text))
            return "\n".join(text_content)

        elif extension == '.pptx':
            prs = pptx.Presentation(file_stream)
            for i, slide in enumerate(prs.slides):
                text_content.append(f"--- Slide {i+1} ---")
                for shape in slide.shapes:
                    if hasattr(shape, "text_frame") and shape.text_frame:
                        text_content.append(shape.text_frame.text)
            return "\n".join(text_content)
        
        elif extension in ['.csv', '.txt', '.md', '.json']:
            # These are text-based, just decode
            return binary_data.decode('utf-8')

        else:
            return f"[Error: Unsupported file type '{extension}'. Could not parse file.]"

    except Exception as e:
        # This catches errors like Base64 decoding
        print(f"[ERROR] Failed to parse file {file_name}: {e}")
        return f"[Error: Could not read file {file_name}. It may be corrupted or an unsupported format.]"

# -----------------------------------------------------------------
# 3. SERVER SETUP (with increased buffer size)
# -----------------------------------------------------------------

# --- MODIFIED: max_http_buffer_size moved here ---
sio = socketio.AsyncServer(
    async_mode="asgi", 
    cors_allowed_origins="*",
    max_http_buffer_size=200 * 1024 * 1024 # 200MB
)
# --- END MODIFIED ---

app = FastAPI()

# --- MODIFIED: Removed the failing 'engineio_options' ---
sio_app = socketio.ASGIApp(
    socketio_server=sio, 
    other_asgi_app=app
)
# --- END MODIFIED ---

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# -----------------------------------------------------------------
# 4. STARTUP LOGIC (Unchanged)
# -----------------------------------------------------------------
@app.on_event("startup")
async def startup_event():
    global compiled_graph, client
    print("Starting MCP Client Server...")
    
    try:
        with open("src/mcp_config.json", "r") as f:
            mcp_config = json.load(f)
        print("Loaded MCP server configuration from src/mcp_config.json")
    except FileNotFoundError:
        print("ERROR: src/mcp_config.json not found.")
        return
    except json.JSONDecodeError:
        print("ERROR: Could not parse src/mcp_config.json.")
        return

    print("Connecting to MCP servers and fetching tools...")
    try:
        client = MultiServerMCPClient(mcp_config["mcpServers"])
        tools = await client.get_tools()
        print(f"Successfully loaded {len(tools)} tools:")
        for tool in tools:
            print(f"- {tool.name}")
    except Exception as e:
        print(f"FATAL: Could not load tools via MultiServerMCPClient: {e}")
        return

    compiled_graph = build_graph(tools)
    print("Chatbot graph compiled successfully.")
    print("\n--- Jarvis is online. Waiting for frontend connection... ---")

# -----------------------------------------------------------------
# 5. SOCKET.IO EVENT HANDLERS (Unchanged)
# -----------------------------------------------------------------

@sio.event
async def connect(sid, environ):
    print(f"[Socket.IO] Frontend connected: {sid}")

@sio.event
async def disconnect(sid):
    print(f"[Socket.IO] Frontend disconnected: {sid}")

# Frontend can request cancellation of in-progress processing for this socket
@sio.event
async def cancel_processing(sid, data):
    task = active_tasks.pop(sid, None)
    if task and not task.done():
        print(f"[{sid}] Received cancel request from frontend. Cancelling task...")
        task.cancel()
    else:
        print(f"[{sid}] Received cancel request but no active task found.")

# Track active processing tasks per socket id so they can be cancelled by the frontend
active_tasks: dict[str, asyncio.Task] = {}

@sio.event
async def chat_message(sid, data):
    """
    Handles incoming chat messages from the frontend.
    """
    if not compiled_graph:
        await sio.emit('ai_response', {'chunk': 'Error: Graph not initialized.'}, to=sid)
        return

    user_input = data.get("message", "").strip()
    file_base64 = data.get("file_base64") 
    file_name = data.get("file_name")
    
    persistent_thread_id = data.get("thread_id", sid) 
    
    if not user_input and not file_base64:
        print("[Request] Received empty message and no file.")
        return 

    final_input = user_input
    
    if file_base64 and file_name:
        print(f"[Context] Receiving file: {file_name}. Parsing...")
        
        try:
            loop = asyncio.get_running_loop()
            extracted_text = await asyncio.wait_for(
                loop.run_in_executor(
                    None, parse_file_content, file_name, file_base64
                ),
                timeout=FILE_PARSE_TIMEOUT
            )
            print(f"[Context] File parsed successfully. ({len(extracted_text)} chars)")
        
        except asyncio.TimeoutError:
            print(f"[ERROR] File parsing timed out for {file_name}")
            extracted_text = f"[Error: File parsing timed out after {FILE_PARSE_TIMEOUT} seconds. The file may be too large or complex.]"
        
        except Exception as e:
            print(f"[ERROR] File parsing failed in executor: {e}")
            extracted_text = f"[Error: Failed to process file {file_name}.]"
        
        context_header = f"--- START OF UPLOADED FILE: {file_name} ---"
        context_footer = f"--- END OF UPLOADED FILE: {file_name} ---"
        
        final_input = (
            f"Please use the following document as context:\n\n"
            f"{context_header}\n"
            f"{extracted_text}\n"
            f"{context_footer}\n\n"
            f"My question is: {user_input if user_input else f'Please analyze the file {file_name}'}"
        )
        
        temp_thread_id = f"oneshot_{sid}_{int(time.time())}"
        run_config = {"configurable": {"thread_id": temp_thread_id}}
        print(f"\n[Request from {sid}] User: {user_input}")
        print(f"[Request from {sid}] File: {file_name}")
        print(f"[Thread] Using temporary 'one-shot' thread: {temp_thread_id}")
    else:
        run_config = {"configurable": {"thread_id": persistent_thread_id}}
        print(f"\n[Request from {sid}] User: {user_input}")
        print(f"[Thread] Using persistent thread: {persistent_thread_id}")
    
    print("Jarvis: ...")

    # Cancel any existing task for this sid (we'll start a fresh one)
    if sid in active_tasks:
        task = active_tasks.pop(sid)
        if not task.done():
            print(f"[{sid}] Cancelling previous task before starting a new one...")
            task.cancel()

    async def _process_and_stream():
        try:
            async for event_type, content in _stream_graph_logic(final_input, compiled_graph, run_config):
                if event_type == "ai":
                    await sio.emit('ai_response', {'chunk': content}, to=sid)
                elif event_type == "debug":
                    print(content.strip())
                elif event_type == "tool_call":
                    await sio.emit('tool_call', {'tool_name': content}, to=sid)

            await sio.emit('ai_response_end', to=sid)
            print("[Jarvis] Response stream complete.")

        except asyncio.CancelledError:
            # Task was cancelled by the frontend
            print(f"[{sid}] Processing cancelled by frontend.")
            try:
                await sio.emit('ai_response_end', to=sid)
            except Exception:
                pass
        except Exception as e:
            print(f"An error occurred: {e}")
            try:
                await sio.emit('ai_response', {'chunk': f'An error occurred: {e}'}, to=sid)
                await sio.emit('ai_response_end', to=sid)
            except Exception:
                pass

    # Start the processing in a background task and store it so it can be cancelled
    task = asyncio.create_task(_process_and_stream())
    active_tasks[sid] = task

# -----------------------------------------------------------------
# 6. SERVER RUNNER (Unchanged)
# -----------------------------------------------------------------
if __name__ == "__main__":
    print("Starting Socket.IO server on http://127.0.0.1:8001")
    uvicorn.run(
        sio_app, 
        host="127.0.0.1", 
        port=8001
    )
